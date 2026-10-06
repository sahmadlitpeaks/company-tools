"""Read lead spreadsheets (CSV or XLSX), map their columns onto CRM fields and
normalize each row, so the preview and the import share one interpretation."""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.models.crm import LEAD_PRIORITIES, LEAD_STATUSES
from app.schemas.crm import normalize_tags

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

# Fields a column can map to. Every field but `notes` takes at most one column.
IMPORT_FIELDS = (
    "name", "email", "phone", "company", "value", "notes", "status", "priority",
    "tags", "next_step", "follow_up_date", "expected_close_date", "lost_reason",
)
_SYNONYMS = {
    "name": ["name", "full name", "contact", "contact name", "lead name", "customer name"],
    "email": ["email", "e mail", "email address", "mail"],
    "phone": ["phone", "mobile", "telephone", "tel", "phone number", "mobile number", "whatsapp"],
    "company": [
        "company", "company name", "organization", "organisation", "account",
        "business", "department / company",
    ],
    "value": ["value", "deal value", "amount", "deal size"],
    "notes": [
        "notes", "note", "message", "full message", "comments", "comment",
        "description", "inquiry", "enquiry",
    ],
    "status": ["status", "stage", "pipeline stage"],
    "priority": ["priority", "rating"],
    "tags": ["tags", "tag", "labels"],
    "next_step": ["next step", "suggested next step", "next action"],
    "follow_up_date": ["follow up", "follow up date", "next follow up"],
    "expected_close_date": ["expected close", "expected close date", "close date"],
    "lost_reason": ["lost reason", "reason lost"],
}
_STATUS_ALIASES = {
    "not contacted": "new", "open": "new", "proposal sent": "proposal",
    "negotiating": "negotiation", "closed won": "won", "closed lost": "lost",
}
_LIMITS = {"name": 255, "email": 320, "phone": 64, "company": 255, "next_step": 255, "lost_reason": 255}
# Characters a spreadsheet may treat as a formula when a cell starts with them.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


class ImportFileError(ValueError):
    """The upload can't be read as a lead table at all."""


@dataclass
class Table:
    sheets: list[str]
    sheet: str | None
    columns: list[str]
    rows: list[list[str]]


@dataclass
class LeadRow:
    row: int  # 1-based spreadsheet row, counting the header
    data: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    error: str | None = None

    @property
    def key(self) -> str | None:
        email = self.data.get("email")
        return email.lower() if email else None


def _header_key(text: str) -> str:
    return " ".join(re.sub(r"[_\-]+", " ", text).lower().split())


_LOOKUP = {_header_key(s): f for f, names in _SYNONYMS.items() for s in names}


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        # Spreadsheets store phone numbers and IDs as numbers.
        return str(int(value))
    return str(value).strip()


def _unique_columns(header: list[str]) -> list[str]:
    columns: list[str] = []
    for index, raw in enumerate(header, 1):
        name = raw or f"Column {index}"
        candidate, n = name, 2
        while candidate in columns:
            candidate, n = f"{name} ({n})", n + 1
        columns.append(candidate)
    return columns


def suggest_mapping(columns: list[str]) -> dict[str, str | None]:
    mapping: dict[str, str | None] = {}
    used: set[str] = set()
    for column in columns:
        target = _LOOKUP.get(_header_key(column))
        if target and (target == "notes" or target not in used):
            mapping[column] = target
            used.add(target)
        else:
            mapping[column] = None
    return mapping


def _table_from_rows(raw_rows, sheets: list[str], sheet: str | None) -> Table:
    rows = [[_cell(v) for v in row] for row in raw_rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        raise ImportFileError("The file has no rows to import.")
    columns = _unique_columns(rows[0])
    body = [(row + [""] * len(columns))[: len(columns)] for row in rows[1:]]
    if len(body) > MAX_ROWS:
        raise ImportFileError(f"Import at most {MAX_ROWS} rows at a time.")
    return Table(sheets=sheets, sheet=sheet, columns=columns, rows=body)


def _read_csv(raw: bytes) -> Table:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return _table_from_rows(csv.reader(io.StringIO(text), dialect), [], None)


def _read_xlsx(raw: bytes, sheet: str | None) -> Table:
    import openpyxl  # imported lazily: only spreadsheet uploads need it

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise ImportFileError("The spreadsheet could not be read. Save it as .xlsx or .csv.") from exc
    try:
        names = list(workbook.sheetnames)
        if sheet and sheet not in names:
            raise ImportFileError(f"The workbook has no sheet named “{sheet}”.")
        if not sheet:
            # Prefer the first sheet whose header row looks like contacts, so a
            # cover/overview sheet at the front doesn't hide the lead list.
            sheet = names[0]
            for name in names:
                first = next(workbook[name].iter_rows(max_row=1, values_only=True), ())
                mapped = set(suggest_mapping([_cell(v) for v in first]).values())
                if {"email", "name"} & mapped:
                    sheet = name
                    break
        rows = workbook[sheet].iter_rows(values_only=True)
        limited = []
        for index, row in enumerate(rows):
            if index > MAX_ROWS + 1:
                raise ImportFileError(f"Import at most {MAX_ROWS} rows at a time.")
            limited.append(row)
        return _table_from_rows(limited, names, sheet)
    finally:
        workbook.close()


def read_table(filename: str, raw: bytes, sheet: str | None = None) -> Table:
    if len(raw) > MAX_BYTES:
        raise ImportFileError("Files can be at most 5 MB.")
    if not raw:
        raise ImportFileError("The file is empty.")
    lower = (filename or "").lower()
    if lower.endswith(".xlsx") or raw[:2] == b"PK":
        return _read_xlsx(raw, sheet)
    if lower.endswith((".xls", ".xlsm", ".numbers", ".ods")):
        raise ImportFileError("Save the spreadsheet as .xlsx or .csv to import it.")
    return _read_csv(raw)


def validate_mapping(columns: list[str], mapping: dict[str, str | None]) -> dict[str, str | None]:
    clean: dict[str, str | None] = {}
    used: set[str] = set()
    for column in columns:
        target = mapping.get(column)
        if target is None:
            clean[column] = None
            continue
        if target not in IMPORT_FIELDS:
            raise ImportFileError(f"“{target}” is not a field leads can import into.")
        if target != "notes" and target in used:
            raise ImportFileError(f"Only one column can fill {target.replace('_', ' ')}.")
        used.add(target)
        clean[column] = target
    if not {"name", "email", "phone"} & used:
        raise ImportFileError("Match at least one column to name, email or phone.")
    return clean


def _unescape(value: str) -> str:
    # Our export prefixes formula-like cells with an apostrophe; undo that.
    if value.startswith("'") and value[1:2] and value[1] in FORMULA_PREFIXES:
        return value[1:]
    return value


def _date(value: str) -> date | None:
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        return None


def normalize_row(row_number: int, values: list[str], columns: list[str], mapping: dict) -> LeadRow:
    lead = LeadRow(row=row_number)
    notes: list[tuple[str, str]] = []
    for column, value in zip(columns, values):
        target = mapping.get(column)
        value = _unescape(value.strip())
        if not target or not value:
            continue
        if target == "notes":
            notes.append((column, value))
        else:
            lead.data[target] = value
    if notes:
        lead.data["notes"] = notes[0][1] if len(notes) == 1 else "\n\n".join(f"{c}: {v}" for c, v in notes)
    data = lead.data
    if not any(data.get(k) for k in ("name", "email", "phone")):
        lead.error = "No name, email or phone"
        return lead

    for key, limit in _LIMITS.items():
        if data.get(key) and len(data[key]) > limit:
            data[key] = data[key][:limit]
            lead.warnings.append(f"{key.replace('_', ' ').capitalize()} was shortened to {limit} characters")
    if data.get("email") and "@" not in data["email"]:
        lead.warnings.append("Email doesn't look like an address")

    if "value" in data:
        cleaned = re.sub(r"[^\d.\-]", "", data["value"])
        try:
            amount = Decimal(cleaned)
            if amount < 0 or amount >= Decimal("1e10"):
                raise InvalidOperation
            data["value"] = amount.quantize(Decimal("0.01"))
        except InvalidOperation:
            lead.warnings.append(f"Value “{data['value']}” isn't a number and was left blank")
            del data["value"]

    status = data.pop("status", None)
    if status:
        normalized = _STATUS_ALIASES.get(status.lower(), status.lower())
        if normalized in LEAD_STATUSES:
            data["status"] = normalized
        else:
            lead.warnings.append(f"Stage “{status}” isn't recognised; imported as new")
    data.setdefault("status", "new")
    if data["status"] == "lost" and not data.get("lost_reason"):
        data["lost_reason"] = "Not recorded (imported)"

    if "priority" in data:
        if data["priority"].lower() in LEAD_PRIORITIES:
            data["priority"] = data["priority"].lower()
        else:
            lead.warnings.append(f"Priority “{data['priority']}” isn't high, medium or low; left blank")
            del data["priority"]

    if "tags" in data:
        data["tags"] = normalize_tags(re.split(r"[,;|]", data["tags"]))

    for key in ("follow_up_date", "expected_close_date"):
        if key in data:
            parsed = _date(data[key])
            if parsed is None:
                lead.warnings.append(f"{key.replace('_', ' ').capitalize()} “{data[key]}” isn't a YYYY-MM-DD date; left blank")
                del data[key]
            else:
                data[key] = parsed
    return lead


def normalize_table(table: Table, mapping: dict) -> list[LeadRow]:
    # Row 1 is the header, so data starts on spreadsheet row 2.
    return [normalize_row(i, values, table.columns, mapping) for i, values in enumerate(table.rows, 2)]
